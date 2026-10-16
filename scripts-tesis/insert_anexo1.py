#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Inserta el ANEXO 1 (Índice del Trabajo Final y contenidos mínimos) al final
del documento existente, SIN borrar ni regenerar nada. Abre el .docx, agrega
los párrafos del anexo respetando el formato del cuerpo (Arial, mismas
jerarquías) y guarda sobre el mismo archivo.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "/Users/fran/Desktop/Tesis/TFG-OFIX-Final-Unificado.docx"
FONT_NAME = "Arial"
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document(PATH)


def _set_run(run, size=12, bold=False, italic=False, color=BLACK, name=FONT_NAME):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def section_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(text)
    _set_run(r, size=16, bold=True)


def para(text, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    _set_run(r, size=12)


def requisito(titulo, requerimiento):
    """Encabezado de requisito: título en negrita + enunciado verbatim en cursiva."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(titulo)
    _set_run(r, size=12, bold=True)
    if requerimiento:
        r2 = p.add_run(" " + requerimiento)
        _set_run(r2, size=12, italic=True)


# ---------------------------------------------------------------------------
# CONTENIDO DEL ANEXO 1
# ---------------------------------------------------------------------------
page_break()
section_title("Anexo 1 — Índice del Trabajo Final y contenidos mínimos")

para(
    "El presente anexo reproduce de manera literal los contenidos mínimos "
    "exigidos por la guía del Trabajo Final de Graduación y establece, para "
    "cada uno de ellos, la correspondencia con la sección del trabajo en la "
    "que es desarrollado. Su propósito es doble: ofrecer al lector una matriz "
    "de trazabilidad que verifique la cobertura íntegra de los requisitos "
    "formales, y sintetizar los principales resultados obtenidos en cada "
    "dimensión del plan de negocio. Los enunciados que anteceden a cada "
    "apartado se transcriben en su redacción original; la descripción "
    "subsiguiente remite a la numeración interna del documento."
)

requisito(
    "Carátula. Resumen Ejecutivo.",
    "Debe ser corto y conciso (una página máxima) y expresar brevemente la "
    "idea general del proyecto, incluyendo: la descripción de la necesidad, "
    "la idea de negocio y el tamaño de la oportunidad; el segmento objetivo; "
    "las ventajas competitivas, el diferenciador del proyecto y su valor "
    "agregado; el equipo emprendedor; y la rentabilidad del negocio, la "
    "inversión requerida y la rentabilidad para el inversor."
)
para(
    "Estos contenidos se desarrollan en la Carátula y el Resumen Ejecutivo "
    "del documento. Este último sintetiza, en extensión acotada, la "
    "oportunidad asociada a la informalidad y la opacidad del mercado de "
    "oficios en Argentina, la propuesta de OFIX como marketplace de "
    "contratación formal, el dimensionamiento de la oportunidad (SAM "
    "subregional de USD 2,9 mil millones y SOM para Argentina de USD 2,85 a "
    "5,70 millones), la composición del equipo fundador y los principales "
    "indicadores de rentabilidad (inversión inicial de USD 87.694, VAN de "
    "USD 82.015 y TIR del 48 %)."
)

requisito("Índice de contenidos. Agradecimientos. Introducción.", "")
para(
    "Se desarrollan, respectivamente, en el Índice de contenidos, los "
    "Agradecimientos y la Introducción. Esta última precisa la motivación y "
    "relevancia del proyecto, el objetivo general y los objetivos "
    "específicos, la estrategia metodológica —articulación de fuentes "
    "primarias (encuesta, entrevistas en profundidad y pruebas reales de "
    "campañas) y secundarias— y la estructura del documento."
)

requisito(
    "I – Evaluación del problema.",
    "Análisis y selección de la oportunidad de negocio digital. Investigación "
    "y validación de la oportunidad, necesidad y/o problema identificado. "
    "Investigación de mercado. Características del público objetivo, "
    "preferencias y necesidades. Árbol del problema y User Persona."
)
para(
    "Estos contenidos se abordan en el Capítulo I. La selección y validación "
    "de la oportunidad se sustenta en la descripción y evidencias del "
    "problema (secciones 1.1 y 1.2), que documentan la informalidad "
    "estructural del sector y la dependencia del boca a boca como canal "
    "dominante de contratación. La caracterización del público objetivo se "
    "desarrolla mediante la identificación de los usuarios afectados (1.3), "
    "el árbol del problema (1.4) y la elaboración de tres User Personas "
    "representativas de los segmentos hogar, pyme gastronómica y trabajador "
    "de oficio (1.5). La investigación de mercado (1.7) integra evidencia "
    "cuantitativa —encuesta con 102 respuestas del lado de la demanda y 21 "
    "del lado de la oferta— y cualitativa —ocho entrevistas en profundidad "
    "(1.8)—, cuyas conclusiones se consolidan en la sección 1.9."
)

requisito(
    "II – Solución Producto Digital.",
    "Propuesta de Valor Única. Desarrollo del MVP. Fundamento y justificación "
    "de su alcance. Modelo de priorización utilizado. ¿Cómo la empresa va a "
    "ganar dinero? Desarrollo profundo del modelo de negocio sugerido. "
    "Business Model Canvas. Tamaño de la oportunidad de negocio (TAM, SAM, "
    "SOM). Roadmap de desarrollo de producto de 3 a 5 años con hitos "
    "esperados para las distintas dimensiones del roadmap. Desarrollo del "
    "stack tecnológico propuesto, incluyendo la infraestructura de servicio "
    "y su fundamento. Desarrollo del marco competitivo exhaustivo y matriz de "
    "posicionamiento que incorpore la solución propuesta en el marco "
    "descripto. Desarrollo UX/ Ui y Prototipado."
)
para(
    "Estos contenidos se desarrollan en el Capítulo II. La propuesta de valor "
    "y el alcance del MVP se fundamentan a partir de la descripción de la "
    "solución (2.1), su deseabilidad (2.2) y un modelo de priorización MoSCoW "
    "(2.5). El modelo de monetización y el modelo de negocio se exponen en "
    "las secciones 2.4 y 2.7 (Business Model Canvas), complementados por el "
    "Value Proposition Canvas (2.8) y el análisis FODA (2.9). La factibilidad "
    "técnica y el stack tecnológico, con su infraestructura de servicio, se "
    "justifican en la sección 2.3. El roadmap de producto a cinco años "
    "(2026–2030) se detalla en la sección 2.6. El marco competitivo se aborda "
    "mediante un benchmark de once plataformas y una entrevista en "
    "profundidad al competidor Reeparando (2.10), y se sintetiza en una "
    "matriz de posicionamiento (2.11). El desarrollo UX/UI y el prototipado "
    "navegable, junto con sus testeos de usabilidad, se presentan en la "
    "sección 2.12. El dimensionamiento de la oportunidad (TAM, SAM y SOM) se "
    "calcula en la sección 2.13, y se complementa con los análisis "
    "estratégicos PESTEL (2.14), de las cinco fuerzas de Porter (2.15), de la "
    "cadena de valor (2.16) y de la tipología de industria según Barney "
    "(2.17)."
)

requisito(
    "III – Pn de Marketing [sic].",
    "Desarrollar los objetivos y la correspondiente planificación de "
    "estrategias y acciones de Marketing Digital. Definir audiencias y "
    "canales y mensajes, contenido, redes, performance. Plan de Medios. "
    "Embudo de conversión."
)
para(
    "Estos contenidos se desarrollan en el Capítulo III. La planificación de "
    "estrategias y acciones de marketing digital, con la definición de "
    "audiencias, canales, mensajes y contenidos, se aborda en las secciones "
    "3.1 a 3.3, segmentada por tipo de usuario. El plan de medios se valida "
    "empíricamente mediante pruebas reales de campañas ejecutadas en Meta Ads "
    "y Google Ads con presupuestos controlados (3.4). La medición de "
    "performance se formaliza a través de los KPI del producto y del negocio "
    "(3.5) y del embudo de conversión —modelo AARRR para los lados de la "
    "demanda y de la oferta, junto con las métricas de salud del marketplace— "
    "(3.6)."
)

requisito(
    "IV – Equipo.",
    "Características de los integrantes del equipo de trabajo que formarán "
    "parte de la construcción y puesta en funcionamiento del producto que se "
    "quiera desarrollar. Habilidades técnicas, habilidades de liderazgo, "
    "habilidades blandas. Organigrama inicial y desarrollo esperado de la "
    "estructura a futuro."
)
para(
    "Estos contenidos se desarrollan en el Capítulo IV. La composición del "
    "equipo durante el primer año de operación, incluyendo el equipo fundador "
    "y las incorporaciones planificadas por trimestre, se expone en la "
    "sección 4.1. El organigrama inicial y su evolución esperada se presentan "
    "en la sección 4.2, y las habilidades técnicas, de liderazgo y blandas, "
    "así como su cobertura progresiva, se analizan en la sección 4.3."
)

requisito(
    "V – Costos, Finanzas e Inversión.",
    "Unit Economics. Principales variables para el desarrollo del negocio. "
    "EERR (P&L). TIR y VAN. Análisis de costos, inversión. Plan de 3 a 5 "
    "años. CAC y LTV."
)
para(
    "Estos contenidos se desarrollan en el Capítulo V. Las principales "
    "variables del negocio se modelan a partir de los costos variables y "
    "fijos (5.2 y 5.3) y de la proyección de ingresos bajo el modelo PxQ "
    "(5.4). El Estado de Resultados (P&L) proyectado se presenta en la "
    "sección 5.6, con ingresos que evolucionan de USD 693.651 en el primer "
    "año a USD 4.012.723 en el quinto. La evaluación financiera (5.7) arroja "
    "un VAN de USD 82.015 y una TIR del 48 %, sustancialmente superior al "
    "WACC del 27 %, con un período de recupero descontado de 44 meses. Los "
    "unit economics, la contribución marginal y el punto de equilibrio —que "
    "incorporan el CAC y el LTV— se analizan en la sección 5.8, mientras que "
    "las NOF, el CAPEX, el flujo de fondos libres y la inversión total "
    "requerida (USD 87.694 iniciales) se desarrollan en las secciones 5.9 y "
    "5.10, bajo los supuestos macroeconómicos explicitados en la sección "
    "5.11."
)

requisito(
    "VI – Aspectos Legales.",
    "Desarrollo del marco regulatorio vigente y esperado para la industria "
    "del producto o servicio propuesto."
)
para(
    "Estos contenidos se desarrollan en el Capítulo VI, que estructura el "
    "marco regulatorio aplicable a una plataforma de intermediación de "
    "servicios en Argentina según una secuencia trimestral: bases legales "
    "(6.1), compliance, seguridad y protección de datos (6.2), consolidación "
    "jurídica y escalabilidad (6.3) y auditoría anual y estándares (6.4)."
)

requisito(
    "VII – Plan de Implementación y Riesgos.",
    "Riesgos potenciales (beneficios y problemas) que pueden afectar la "
    "construcción y puesta en marcha del producto digital. Plan de "
    "implementación congruente a nivel de tiempos, costos, recursos y "
    "factibilidad técnica."
)
para(
    "Estos contenidos se desarrollan en el Capítulo VII. El plan de "
    "implementación del primer año se organiza por áreas funcionales "
    "—producto, operaciones, marketing y administración y finanzas— con hitos "
    "trimestrales (secciones 7.1 a 7.5). El análisis de riesgos potenciales "
    "se presenta en la sección 7.6, y la proyección de escalabilidad regional "
    "a cinco años, con la evaluación comparativa de los mercados de Uruguay, "
    "México, Colombia y Brasil, en la sección 7.7."
)

doc.save(PATH)
print("OK -> Anexo 1 insertado en", PATH)
print("Total de párrafos ahora:", len(doc.paragraphs))
